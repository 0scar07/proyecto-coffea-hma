# -*- coding: utf-8 -*-
"""Convierte las 4 tablas CSV de docs/validacion_externa/datos/ a .xlsx formateados
para lectura humana, en docs/validacion_externa/datos/excel/. NO cambia ningun valor:
los CSV originales (que el resto del pipeline sigue leyendo) no se tocan; el redondeo
y las columnas Si/No son solo de presentacion en el .xlsx.

Uso: python docs/validacion_externa/fuente_pdf/generar_excel.py
"""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.dirname(AQUI)  # docs/validacion_externa/
DATOS_DIR = os.path.join(OUT_DIR, "datos")
EXCEL_DIR = os.path.join(DATOS_DIR, "excel")
os.makedirs(EXCEL_DIR, exist_ok=True)

TERRACOTA = "A8432B"
CREMA = "F1ECE3"
CREMA_CLARO = "F8F5F0"
BLANCO = "FFFFFF"
TINTA = "201C18"

# Decimales de presentacion por nombre de columna (no cambia el CSV, solo el .xlsx)
DECIMALES = {
    "r2": 4, "mi_r2": 4, "paper_r2": 4, "mi_referencia_aprox_r2": 4,
    "rmse": 4, "mae": 4, "k": 4, "r_o_k": 4,
    "mi_b_1_sobre_k": 2, "paper_b": 2,
    "P0_o_K": 2, "K": 1, "Ti": 1, "K_sobre_max_obs": 2,
    "mi_K_a": 1, "mi_Ti_x0": 1, "paper_a_K": 1, "paper_x0_Ti": 1,
    "mi_referencia_aprox_K": 1, "mi_referencia_aprox_x0": 0,
    "mat": 0, "media_control_Ni": 2, "media_tratamiento": 2, "media_control": 2,
    "incremento_pct": 1, "n_escenario": 0, "sd_control_usada": 2,
    "sd_tratamiento_usada": 2, "t_welch": 3, "p_welch": 4,
    "n_fechas": 0, "n_params": 0, "grados_libertad": 0,
}
COLUMNAS_BOOLEANAS = {"dibuja_K_criterio_app", "significativo_p<0.05"}

LEEME = {
    "resultados_curvas": {
        "que_contiene": "Ajuste de los 3 modelos de la app (Exponencial, Logístico, Gompertz) "
                         "por variable, para las fuentes León-Burgos et al. (2022) y Siqueira et al. (1998).",
        "fuente": "León-Burgos et al. (2022) y Siqueira et al. (1998) — ver la columna 'fuente' en cada fila.",
        "como_leer": "Cada fila es un modelo aplicado a una variable (y, en Siqueira, a un tratamiento). "
                      "'estado' dice si convergió; K/k/Ti solo aplican a Logístico y Gompertz; "
                      "'dibuja_K_criterio_app' indica si el criterio de plausibilidad de la app permite mostrar la asíntota.",
    },
    "comparacion_publicado": {
        "que_contiene": "Mi modelo Logístico convertido a la parametrización sigmoidal del paper, "
                         "comparado con los parámetros publicados y con mi referencia aproximada previa.",
        "fuente": "León-Burgos et al. (2022).",
        "como_leer": "Una fila por variable (área foliar, número de hojas, altura). Compare las "
                      "columnas 'mi_*' contra 'paper_*': diferencias de 1-5% indican que mi "
                      "implementación reproduce el ajuste publicado.",
    },
    "efecto": {
        "que_contiene": "Porcentaje de incremento de cada tratamiento/consorcio sobre su control, "
                         "en el tiempo (Siqueira) o por variable (Vallejos-Torres).",
        "fuente": "Siqueira et al. (1998) y Vallejos-Torres et al. (2021) — ver la columna 'fuente'.",
        "como_leer": "'incremento_pct' es el % de la media del tratamiento sobre la media del "
                      "control en esa fecha/variable. Las filas de Siqueira tienen 'mat' (meses); "
                      "las de Vallejos-Torres, no.",
    },
    "sensibilidad_vallejos_torres": {
        "que_contiene": "Prueba t de Welch de cada consorcio contra el control, bajo 4 escenarios "
                         "explícitos: n=18 o n=108, y el valor entre paréntesis del paper interpretado "
                         "como desviación estándar o como error estándar.",
        "fuente": "Vallejos-Torres et al. (2021).",
        "como_leer": "Cada tratamiento/variable aparece 4 veces (una por escenario). "
                      "'significativo_p<0.05' cambia casi siempre según 'interpretacion_parentesis', "
                      "no según 'n_escenario' — ver informe.md/informe.pdf para la explicación.",
    },
}


def formatear_valor(nombre_col, valor):
    if pd.isna(valor):
        return ""
    if nombre_col in COLUMNAS_BOOLEANAS:
        return "Sí" if bool(valor) else "No"
    if nombre_col in DECIMALES:
        dec = DECIMALES[nombre_col]
        try:
            v = round(float(valor), dec)
            return int(v) if dec == 0 else v
        except (TypeError, ValueError):
            return valor
    return valor


def autoajustar_columnas(ws, df):
    for idx, col in enumerate(df.columns, start=1):
        letra = get_column_letter(idx)
        max_len = max(
            [len(str(col))] + [len(str(formatear_valor(col, v))) for v in df[col]]
        )
        ws.column_dimensions[letra].width = min(max(max_len + 2, 10), 45)


def escribir_hoja_leeme(wb, nombre_csv, n_filas, n_cols):
    ws = wb.active
    ws.title = "Léeme"
    info = LEEME[nombre_csv]
    filas = [
        (f"{nombre_csv}.xlsx", None),
        ("", None),
        ("Qué contiene", info["que_contiene"]),
        ("De qué fuente sale", info["fuente"]),
        ("Cómo leerlo", info["como_leer"]),
        ("Tamaño", f"{n_filas} filas × {n_cols} columnas"),
        ("Nota", "Los valores se redondearon para lectura humana; el CSV de origen en "
                 "datos/ conserva la precisión completa. Ningún dato fue alterado."),
    ]
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 90
    for i, (etiqueta, texto) in enumerate(filas, start=1):
        c1 = ws.cell(row=i, column=1, value=etiqueta)
        c1.font = Font(bold=True, color=TERRACOTA, name="Calibri", size=11)
        if texto is not None:
            c2 = ws.cell(row=i, column=2, value=texto)
            c2.alignment = Alignment(wrap_text=True, vertical="top")
            c2.font = Font(name="Calibri", size=10.5)
    ws.row_dimensions[1].height = 22
    ws.cell(row=1, column=1).font = Font(bold=True, color=TERRACOTA, name="Georgia", size=14)


def escribir_hoja_datos(wb, nombre_csv, df):
    ws = wb.create_sheet(title=nombre_csv[:31])
    encabezado_fill = PatternFill("solid", fgColor=TERRACOTA)
    encabezado_font = Font(bold=True, color=BLANCO, name="Calibri", size=10.5)
    fill_par = PatternFill("solid", fgColor=BLANCO)
    fill_impar = PatternFill("solid", fgColor=CREMA_CLARO)

    for j, col in enumerate(df.columns, start=1):
        celda = ws.cell(row=1, column=j, value=col)
        celda.fill = encabezado_fill
        celda.font = encabezado_font
        celda.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")

    for i, (_, fila) in enumerate(df.iterrows(), start=2):
        for j, col in enumerate(df.columns, start=1):
            valor = formatear_valor(col, fila[col])
            celda = ws.cell(row=i, column=j, value=valor)
            celda.fill = fill_par if i % 2 == 0 else fill_impar
            celda.font = Font(name="Calibri", size=10, color=TINTA)

    ws.freeze_panes = "A2"
    ultima_col = get_column_letter(len(df.columns))
    ws.auto_filter.ref = f"A1:{ultima_col}{len(df) + 1}"
    ws.row_dimensions[1].height = 30
    autoajustar_columnas(ws, df)


def generar_xlsx(nombre_csv):
    ruta_csv = os.path.join(DATOS_DIR, f"{nombre_csv}.csv")
    df = pd.read_csv(ruta_csv)
    wb = Workbook()
    escribir_hoja_leeme(wb, nombre_csv, len(df), len(df.columns))
    escribir_hoja_datos(wb, nombre_csv, df)
    ruta_xlsx = os.path.join(EXCEL_DIR, f"{nombre_csv}.xlsx")
    wb.save(ruta_xlsx)
    print(f"{nombre_csv}.xlsx: {os.path.getsize(ruta_xlsx)} bytes, hojas={wb.sheetnames}")
    return ruta_xlsx


if __name__ == "__main__":
    for nombre in ("resultados_curvas", "comparacion_publicado", "efecto", "sensibilidad_vallejos_torres"):
        generar_xlsx(nombre)
    print("\nListo.")
